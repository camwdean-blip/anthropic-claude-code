"""
Logs every processed transaction to an Excel tracking spreadsheet.
"""

import os
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment

import config

COLUMNS = [
    "Timestamp",
    "Property Address",
    "City",
    "State",
    "ZIP",
    "Sale Price",
    "Buyer",
    "Seller",
    "Attorney",
    "Attorney Email",
    "Closing Date",
    "Commission Rate",
    "Commission Amount",
    "Invoice #",
    "TD Sheet File",
    "Invoice File",
    "Email Status",
    "Agent",
    "MLS #",
]


def _init_workbook(path):
    """Create a new tracking workbook with styled headers."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Transaction Log"

    header_font = Font(bold=True, color="FFFFFF", size=11)
    header_fill = PatternFill(start_color="1A3C6E", end_color="1A3C6E", fill_type="solid")

    for col_idx, header in enumerate(COLUMNS, start=1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center")
        ws.column_dimensions[cell.column_letter].width = max(len(header) + 4, 14)

    ws.freeze_panes = "A2"
    os.makedirs(os.path.dirname(path), exist_ok=True)
    wb.save(path)
    return wb


def log_transaction(transaction, invoice_number, commission, td_sheet_path,
                    invoice_path, email_status):
    """
    Append a row to the tracking spreadsheet.

    Args:
        transaction: dict with keys from the CSV row.
        invoice_number: str invoice ID.
        commission: float computed commission amount.
        td_sheet_path: path to generated TD sheet.
        invoice_path: path to generated invoice.
        email_status: str status from the email sender.

    Returns:
        str: path to the tracking spreadsheet.
    """
    path = config.TRACKING_SPREADSHEET

    if os.path.exists(path):
        wb = load_workbook(path)
    else:
        wb = _init_workbook(path)

    ws = wb.active

    row = [
        datetime.now().strftime("%Y-%m-%d %H:%M:%S"),
        transaction["property_address"],
        transaction["city"],
        transaction["state"],
        transaction["zip"],
        float(transaction["sale_price"]),
        transaction["buyer_name"],
        transaction["seller_name"],
        transaction["attorney_name"],
        transaction["attorney_email"],
        transaction["closing_date"],
        float(transaction["commission_rate"]),
        commission,
        invoice_number,
        os.path.basename(td_sheet_path),
        os.path.basename(invoice_path),
        email_status,
        transaction["agent_name"],
        transaction["mls_number"],
    ]

    ws.append(row)
    wb.save(path)
    return path
