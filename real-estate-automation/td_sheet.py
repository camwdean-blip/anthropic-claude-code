"""
Generates a Transaction Detail (TD) sheet as a Word document.
"""

from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import os

import config


def _add_heading_block(doc, title):
    """Add a styled section heading."""
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(title)
    run.bold = True
    run.font.size = Pt(13)
    run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)
    paragraph.space_after = Pt(4)
    return paragraph


def _add_detail_table(doc, rows):
    """Add a two-column detail table with label/value pairs."""
    table = doc.add_table(rows=len(rows), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"

    for i, (label, value) in enumerate(rows):
        label_cell = table.cell(i, 0)
        value_cell = table.cell(i, 1)

        label_cell.text = label
        value_cell.text = str(value)

        for paragraph in label_cell.paragraphs:
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(10)
        for paragraph in value_cell.paragraphs:
            for run in paragraph.runs:
                run.font.size = Pt(10)

        label_cell.width = Inches(2.5)
        value_cell.width = Inches(4.0)

    doc.add_paragraph()  # spacer
    return table


def generate_td_sheet(transaction, output_path=None):
    """
    Generate a Transaction Detail sheet Word document.

    Args:
        transaction: dict with keys from the CSV row.
        output_path: optional explicit path; otherwise auto-generated.

    Returns:
        str: path to the generated .docx file.
    """
    doc = Document()

    # Title
    title = doc.add_heading("Transaction Detail Sheet", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    for run in title.runs:
        run.font.color.rgb = RGBColor(0x1A, 0x3C, 0x6E)

    # MLS / Closing info bar
    doc.add_paragraph(
        f"MLS #: {transaction['mls_number']}    |    "
        f"Closing Date: {transaction['closing_date']}"
    ).alignment = WD_ALIGN_PARAGRAPH.CENTER

    doc.add_paragraph()  # spacer

    # --- Property Information ---
    _add_heading_block(doc, "Property Information")
    _add_detail_table(doc, [
        ("Street Address", transaction["property_address"]),
        ("City", transaction["city"]),
        ("State", transaction["state"]),
        ("ZIP Code", transaction["zip"]),
        ("Sale Price", f"${float(transaction['sale_price']):,.2f}"),
    ])

    # --- Buyer Information ---
    _add_heading_block(doc, "Buyer Information")
    _add_detail_table(doc, [
        ("Buyer Name", transaction["buyer_name"]),
        ("Buyer Email", transaction["buyer_email"]),
    ])

    # --- Seller Information ---
    _add_heading_block(doc, "Seller Information")
    _add_detail_table(doc, [
        ("Seller Name", transaction["seller_name"]),
        ("Seller Email", transaction["seller_email"]),
    ])

    # --- Agent / Brokerage ---
    _add_heading_block(doc, "Agent & Brokerage")
    _add_detail_table(doc, [
        ("Agent Name", transaction["agent_name"]),
        ("License #", transaction["agent_license"]),
        ("Brokerage", transaction["brokerage_name"]),
        ("Brokerage Address", transaction["brokerage_address"]),
    ])

    # --- Attorney ---
    _add_heading_block(doc, "Closing Attorney")
    _add_detail_table(doc, [
        ("Attorney Name", transaction["attorney_name"]),
        ("Attorney Email", transaction["attorney_email"]),
    ])

    # --- Commission ---
    sale_price = float(transaction["sale_price"])
    rate = float(transaction["commission_rate"])
    commission = sale_price * (rate / 100)

    _add_heading_block(doc, "Commission Summary")
    _add_detail_table(doc, [
        ("Commission Rate", f"{rate}%"),
        ("Commission Amount", f"${commission:,.2f}"),
    ])

    # Determine output path
    if output_path is None:
        safe_address = transaction["property_address"].replace(" ", "_").replace("/", "-")
        filename = f"TD_Sheet_{safe_address}.docx"
        output_path = os.path.join(config.OUTPUT_DIR, filename)

    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    doc.save(output_path)
    return output_path
